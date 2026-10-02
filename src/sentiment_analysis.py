#!/usr/bin/env python3
"""
Sentiment Analysis — Sentiment140 Dataset
==========================================
Preprocessing  : lowercase, emoji removal, lemmatization, tokenization,
                 stop-word removal (negations preserved)
Vectorizers    : Count Vectorizer (BoW), Hashing Vectorizer, Word2Vec (avg)
Models         : KNN (k=5), SVM (LinearSVC), LSTM (trainable & W2V embeddings)
Evaluation     : 10-fold stratified cross-validation
Metrics        : Accuracy, Precision, Recall (weighted)
Output         : results/sentiment_analysis_results.xlsx  +  5 PNG plots

Runtime note   : ~1–3 hours on CPU (LSTM × 10 folds dominates).
                 Set CFG['RUN_LSTM'] = False to skip LSTM for a fast run.
"""

import os, re, sys, time, zipfile, warnings, urllib.request
from pathlib import Path

import numpy as np
import pandas as pd

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import seaborn as sns

import nltk
import emoji

from nltk.tokenize import word_tokenize
from nltk.corpus import stopwords
from nltk.stem import WordNetLemmatizer

from sklearn.feature_extraction.text import CountVectorizer, HashingVectorizer
from sklearn.neighbors import KNeighborsClassifier
from sklearn.svm import LinearSVC
from sklearn.model_selection import StratifiedKFold
from sklearn.metrics import accuracy_score, precision_score, recall_score
from sklearn.decomposition import TruncatedSVD

from gensim.models import Word2Vec

import tensorflow as tf
from tensorflow.keras.models import Sequential
from tensorflow.keras.layers import Embedding, LSTM, Bidirectional, Dense, Dropout, SpatialDropout1D
from tensorflow.keras.preprocessing.text import Tokenizer as KerasTokenizer
from tensorflow.keras.preprocessing.sequence import pad_sequences
from tensorflow.keras.callbacks import EarlyStopping, ReduceLROnPlateau
from tensorflow.keras.optimizers import Adam

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

warnings.filterwarnings('ignore')
os.environ['TF_CPP_MIN_LOG_LEVEL'] = '3'

# ════════════════════════════════════════════════════════════════════
#  CONFIGURATION
# ════════════════════════════════════════════════════════════════════
CFG = {
    # ── Data ──────────────────────────────────────────────────────
    'SAMPLE_SIZE':        50_000,       # tweets per class (positive + negative)
    'RANDOM_STATE':       42,
    'DATA_FILE':          'training.1600000.processed.noemoticon.csv',
    'DATA_URL':           'http://cs.stanford.edu/people/alecmgo/trainingandtestdata.zip',

    # ── Cross-validation ──────────────────────────────────────────
    'N_FOLDS':            10,

    # ── KNN ───────────────────────────────────────────────────────
    'KNN_K':              5,
    'KNN_LSA_DIMS':       150,          # TruncatedSVD dims before KNN (sparse → dense)

    # ── Count Vectorizer ──────────────────────────────────────────
    'COUNT_MAX_FEATURES': 50_000,
    'COUNT_MIN_DF':       2,
    'COUNT_NGRAM':        (1, 2),

    # ── Hashing Vectorizer ────────────────────────────────────────
    'HASH_N_FEATURES':    2 ** 17,      # 131 072 buckets
    'HASH_NGRAM':         (1, 2),

    # ── Word2Vec ──────────────────────────────────────────────────
    'W2V_DIM':            100,
    'W2V_WINDOW':         5,
    'W2V_MIN_COUNT':      2,
    'W2V_EPOCHS':         5,
    'W2V_WORKERS':        4,

    # ── LSTM ──────────────────────────────────────────────────────
    'LSTM_MAX_WORDS':     30_000,
    'LSTM_MAX_LEN':       50,
    'LSTM_EMBED_DIM':     128,          # for trainable-embedding variant
    'LSTM_UNITS':         128,
    'LSTM_SPATIAL_DROP':  0.3,
    'LSTM_DENSE_DROP':    0.4,
    'LSTM_EPOCHS':        20,
    'LSTM_BATCH_SIZE':    128,
    'LSTM_PATIENCE':      5,            # EarlyStopping patience
    'LSTM_LR':            1e-3,
    'LSTM_CLIPNORM':      1.0,

    # ── Switches ──────────────────────────────────────────────────
    'RUN_LSTM':           True,         # set False to skip LSTM (saves hours)

    # ── Output ────────────────────────────────────────────────────
    'RESULTS_DIR':        Path('results'),
}

CFG['RESULTS_DIR'].mkdir(exist_ok=True)

# ════════════════════════════════════════════════════════════════════
#  NLTK RESOURCES
# ════════════════════════════════════════════════════════════════════
for _r in ['punkt', 'punkt_tab', 'stopwords', 'wordnet', 'omw-1.4']:
    nltk.download(_r, quiet=True)

# Negation words are excluded from the stop-word list
_NEGATIONS = {
    'no', 'not', 'never', 'neither', 'nor', 'nothing', 'nowhere',
    'nobody', 'none', 'cannot', "can't", "won't", "don't", "doesn't",
    "didn't", "isn't", "aren't", "wasn't", "weren't", "hasn't", "haven't",
    "hadn't", "wouldn't", "shouldn't", "couldn't", "mightn't", "mustn't",
    "needn't", "shan't", "n't",
}
STOP_WORDS = set(stopwords.words('english')) - _NEGATIONS
LEMMATIZER = WordNetLemmatizer()


# ════════════════════════════════════════════════════════════════════
#  DATASET DOWNLOAD & LOADING
# ════════════════════════════════════════════════════════════════════
def _dl_progress(block, block_size, total):
    pct = min(block * block_size / total * 100, 100)
    bar = '█' * int(pct / 2)
    print(f'\r  [{bar:<50}] {pct:.1f}%', end='', flush=True)


def download_dataset() -> None:
    """Download and extract Sentiment140 if not already present."""
    if Path(CFG['DATA_FILE']).exists():
        print(f"  Dataset found: {CFG['DATA_FILE']}")
        return

    zip_path = 'trainingandtestdata.zip'
    if not Path(zip_path).exists():
        print('  Downloading Sentiment140 dataset (~80 MB) …')
        try:
            urllib.request.urlretrieve(CFG['DATA_URL'], zip_path, _dl_progress)
            print()
        except Exception as exc:
            print(f'\n  Auto-download failed: {exc}')
            print('  Manual download: https://www.kaggle.com/datasets/kazanova/sentiment140')
            print(f"  Extract '{CFG['DATA_FILE']}' to the current directory and re-run.")
            sys.exit(1)

    print('  Extracting …')
    with zipfile.ZipFile(zip_path, 'r') as zf:
        zf.extractall('.')
    print('  Dataset ready.')


def load_data() -> tuple:
    """Load 50 K negative + 50 K positive tweets from Sentiment140."""
    print('\n[1/6] Loading dataset …')

    df = pd.read_csv(
        CFG['DATA_FILE'],
        encoding='latin-1',
        header=None,
        names=['polarity', 'id', 'date', 'query', 'user', 'text'],
        usecols=['polarity', 'text'],
    )

    n = CFG['SAMPLE_SIZE']
    neg = df[df['polarity'] == 0].sample(n, random_state=CFG['RANDOM_STATE'])
    pos = df[df['polarity'] == 4].sample(n, random_state=CFG['RANDOM_STATE'])

    data = (
        pd.concat([neg, pos])
        .sample(frac=1, random_state=CFG['RANDOM_STATE'])
        .reset_index(drop=True)
    )

    labels = (data['polarity'] == 4).astype(int).values   # 0 = negative, 1 = positive
    texts  = data['text'].values

    print(f'  {len(data):,} tweets  |  {(labels == 0).sum():,} negative  |  {(labels == 1).sum():,} positive')
    return texts, labels


# ════════════════════════════════════════════════════════════════════
#  PREPROCESSING
# ════════════════════════════════════════════════════════════════════
def _clean(text: str) -> str:
    """Clean raw tweet text."""
    text = text.lower()
    text = re.sub(r'https?://\S+|www\.\S+', '', text)    # URLs
    text = re.sub(r'@\w+', '', text)                      # @mentions
    text = re.sub(r'#\w+', '', text)                      # #hashtags
    text = emoji.replace_emoji(text, replace='')          # emojis
    text = re.sub(r'&\w+;', '', text)                     # HTML entities
    text = re.sub(r'[^a-z\s]', '', text)                  # non-alpha chars
    return re.sub(r'\s+', ' ', text).strip()


def _preprocess(text: str) -> list:
    """Tokenize, lemmatize, and filter stop words (negations kept)."""
    tokens = word_tokenize(_clean(text))
    return [
        LEMMATIZER.lemmatize(t)
        for t in tokens
        if len(t) > 1 and t not in STOP_WORDS
    ]


def preprocess_all(texts) -> tuple:
    """Apply preprocessing pipeline to all tweets; return tokens + joined strings."""
    print('\n[2/6] Preprocessing tweets …')
    t0 = time.time()
    tokenized = []

    for i, text in enumerate(texts):
        tokenized.append(_preprocess(text))
        if (i + 1) % 20_000 == 0:
            elapsed = time.time() - t0
            eta     = elapsed / (i + 1) * (len(texts) - i - 1)
            print(f'  {i+1:>7,}/{len(texts):,}  |  {elapsed:.0f}s elapsed  |  {eta:.0f}s remaining')

    corpus_str = [' '.join(tok) for tok in tokenized]
    print(f'  Done in {time.time() - t0:.1f}s  |  avg tokens/tweet: '
          f'{np.mean([len(t) for t in tokenized]):.1f}')
    return tokenized, corpus_str


# ════════════════════════════════════════════════════════════════════
#  WORD2VEC
# ════════════════════════════════════════════════════════════════════
def train_word2vec(tokenized: list) -> Word2Vec:
    """Train Word2Vec on the full (preprocessed) corpus.
    Unsupervised — labels are not used, so training on all data is acceptable.
    """
    print('\n[3/6] Training Word2Vec embeddings …')
    t0 = time.time()
    model = Word2Vec(
        sentences=tokenized,
        vector_size=CFG['W2V_DIM'],
        window=CFG['W2V_WINDOW'],
        min_count=CFG['W2V_MIN_COUNT'],
        workers=CFG['W2V_WORKERS'],
        epochs=CFG['W2V_EPOCHS'],
        seed=CFG['RANDOM_STATE'],
    )
    print(f'  Vocabulary: {len(model.wv):,} words  |  Done in {time.time() - t0:.1f}s')
    return model


def _doc_to_w2v(tokens: list, wv) -> np.ndarray:
    """Average Word2Vec vectors for a single document."""
    vecs = [wv[t] for t in tokens if t in wv]
    return np.mean(vecs, axis=0).astype(np.float32) if vecs else np.zeros(CFG['W2V_DIM'], dtype=np.float32)


def build_w2v_matrix(tokenized: list, model: Word2Vec) -> np.ndarray:
    return np.array([_doc_to_w2v(tok, model.wv) for tok in tokenized])


def build_lstm_embedding_matrix(keras_tokenizer: KerasTokenizer, w2v: Word2Vec) -> tuple:
    """Create an embedding weight matrix aligned to the Keras tokenizer vocabulary."""
    vocab_size = min(CFG['LSTM_MAX_WORDS'], len(keras_tokenizer.word_index)) + 1
    matrix = np.zeros((vocab_size, CFG['W2V_DIM']), dtype=np.float32)
    hits = 0
    for word, idx in keras_tokenizer.word_index.items():
        if idx >= vocab_size:
            continue
        if word in w2v.wv:
            matrix[idx] = w2v.wv[word]
            hits += 1
    return matrix, vocab_size, hits


# ════════════════════════════════════════════════════════════════════
#  CV HELPERS
# ════════════════════════════════════════════════════════════════════
def _metrics(y_true, y_pred) -> tuple:
    acc = accuracy_score(y_true, y_pred)
    pre = precision_score(y_true, y_pred, average='weighted', zero_division=0)
    rec = recall_score(y_true, y_pred, average='weighted', zero_division=0)
    return acc, pre, rec


def _make_record(model_name: str, vec_name: str, folds: list) -> dict:
    accs  = [f[0] for f in folds]
    precs = [f[1] for f in folds]
    recs  = [f[2] for f in folds]
    return {
        'Model': model_name, 'Vectorizer': vec_name,
        'Accuracy_mean':   float(np.mean(accs)),   'Accuracy_std':   float(np.std(accs)),
        'Precision_mean':  float(np.mean(precs)),  'Precision_std':  float(np.std(precs)),
        'Recall_mean':     float(np.mean(recs)),   'Recall_std':     float(np.std(recs)),
        'Accuracy_folds':  accs,
        'Precision_folds': precs,
        'Recall_folds':    recs,
    }


# ════════════════════════════════════════════════════════════════════
#  CLASSICAL MODEL CV  (KNN / SVM)
# ════════════════════════════════════════════════════════════════════
def run_classical_cv(
    model_name: str,
    clf_factory,
    vec_name: str,
    corpus_arr: np.ndarray,   # numpy array of joined strings
    X_w2v: np.ndarray,
    labels: np.ndarray,
    skf: StratifiedKFold,
) -> dict:
    """10-fold CV for KNN or SVM. CountVectorizer is fit inside each fold (no leakage)."""
    folds = []

    for fold, (tr, va) in enumerate(skf.split(np.zeros(len(labels)), labels), 1):
        y_tr, y_va = labels[tr], labels[va]

        if vec_name == 'Count Vectorizer':
            vec  = CountVectorizer(
                max_features=CFG['COUNT_MAX_FEATURES'],
                min_df=CFG['COUNT_MIN_DF'],
                ngram_range=CFG['COUNT_NGRAM'],
            )
            X_tr = vec.fit_transform(corpus_arr[tr])
            X_va = vec.transform(corpus_arr[va])
            if model_name == 'KNN':
                # Reduce sparse dimensionality before cosine-NN search
                lsa  = TruncatedSVD(n_components=CFG['KNN_LSA_DIMS'], random_state=CFG['RANDOM_STATE'])
                X_tr = lsa.fit_transform(X_tr)
                X_va = lsa.transform(X_va)

        elif vec_name == 'Hashing Vectorizer':
            vec  = HashingVectorizer(
                n_features=CFG['HASH_N_FEATURES'],
                alternate_sign=False,
                ngram_range=CFG['HASH_NGRAM'],
            )
            X_tr = vec.transform(corpus_arr[tr])   # no fit needed — stateless
            X_va = vec.transform(corpus_arr[va])
            if model_name == 'KNN':
                lsa  = TruncatedSVD(n_components=CFG['KNN_LSA_DIMS'], random_state=CFG['RANDOM_STATE'])
                X_tr = lsa.fit_transform(X_tr)
                X_va = lsa.transform(X_va)

        elif vec_name == 'Word2Vec':
            X_tr = X_w2v[tr]
            X_va = X_w2v[va]

        else:
            raise ValueError(f'Unknown vectorizer: {vec_name}')

        clf = clf_factory()
        clf.fit(X_tr, y_tr)
        acc, pre, rec = _metrics(y_va, clf.predict(X_va))
        folds.append((acc, pre, rec))
        print(f'    Fold {fold:2d}/10  |  Acc {acc:.4f}  Pre {pre:.4f}  Rec {rec:.4f}')

    return _make_record(model_name, vec_name, folds)


# ════════════════════════════════════════════════════════════════════
#  LSTM CV
# ════════════════════════════════════════════════════════════════════
def run_lstm_cv(
    tokenized: list,
    labels: np.ndarray,
    skf: StratifiedKFold,
    w2v: Word2Vec,
    variant: str,             # 'trainable' | 'w2v'
) -> dict:
    """10-fold CV for LSTM. Keras tokenizer is fit on the training fold only."""
    vec_name = 'Trainable Embeddings' if variant == 'trainable' else 'Word2Vec Embeddings'
    folds    = []

    for fold, (tr, va) in enumerate(skf.split(np.zeros(len(labels)), labels), 1):
        y_tr, y_va = labels[tr], labels[va]

        train_texts = [' '.join(tokenized[i]) for i in tr]
        val_texts   = [' '.join(tokenized[i]) for i in va]

        # Fit tokenizer on training fold only
        kt = KerasTokenizer(num_words=CFG['LSTM_MAX_WORDS'], oov_token='<OOV>')
        kt.fit_on_texts(train_texts)

        X_tr = pad_sequences(
            kt.texts_to_sequences(train_texts),
            maxlen=CFG['LSTM_MAX_LEN'], padding='post', truncating='post',
        )
        X_va = pad_sequences(
            kt.texts_to_sequences(val_texts),
            maxlen=CFG['LSTM_MAX_LEN'], padding='post', truncating='post',
        )

        # Embedding layer
        if variant == 'trainable':
            embed_layer = Embedding(
                CFG['LSTM_MAX_WORDS'] + 1,
                CFG['LSTM_EMBED_DIM'],
                input_length=CFG['LSTM_MAX_LEN'],
                embeddings_initializer='glorot_uniform',
            )
        else:
            emb_matrix, vocab_size, hits = build_lstm_embedding_matrix(kt, w2v)
            embed_layer = Embedding(
                vocab_size, CFG['W2V_DIM'],
                weights=[emb_matrix],
                input_length=CFG['LSTM_MAX_LEN'],
                trainable=False,
            )

        model = Sequential([
            embed_layer,
            SpatialDropout1D(CFG['LSTM_SPATIAL_DROP']),
            Bidirectional(LSTM(CFG['LSTM_UNITS'], dropout=0.2)),
            Dense(64, activation='relu'),
            Dropout(CFG['LSTM_DENSE_DROP']),
            Dense(1, activation='sigmoid'),
        ])
        model.compile(
            optimizer=Adam(learning_rate=CFG['LSTM_LR'], clipnorm=CFG['LSTM_CLIPNORM']),
            loss='binary_crossentropy', metrics=['accuracy'],
        )

        es = EarlyStopping(
            monitor='val_loss',
            patience=CFG['LSTM_PATIENCE'],
            restore_best_weights=True,
            verbose=0,
        )
        rlr = ReduceLROnPlateau(monitor='val_loss', factor=0.5, patience=2, verbose=0)
        model.fit(
            X_tr, y_tr,
            epochs=CFG['LSTM_EPOCHS'],
            batch_size=CFG['LSTM_BATCH_SIZE'],
            validation_data=(X_va, y_va),
            callbacks=[es, rlr],
            verbose=1,
        )

        y_pred = (model.predict(X_va, verbose=0).flatten() >= 0.5).astype(int)
        acc, pre, rec = _metrics(y_va, y_pred)
        folds.append((acc, pre, rec))
        print(f'    Fold {fold:2d}/10  |  Acc {acc:.4f}  Pre {pre:.4f}  Rec {rec:.4f}')

        tf.keras.backend.clear_session()

    return _make_record('LSTM', vec_name, folds)


# ════════════════════════════════════════════════════════════════════
#  ORCHESTRATION
# ════════════════════════════════════════════════════════════════════
def train_all(tokenized: list, corpus_str: list, labels: np.ndarray, w2v: Word2Vec) -> list:
    print('\n[4/6] Training all models with 10-fold stratified cross-validation …')
    print('      (KNN + LSA for sparse matrices | negation-aware stop words | W2V avg for KNN/SVM)\n')

    skf       = StratifiedKFold(n_splits=CFG['N_FOLDS'], shuffle=True, random_state=CFG['RANDOM_STATE'])
    corpus_np = np.array(corpus_str)
    X_w2v     = build_w2v_matrix(tokenized, w2v)

    knn_fn = lambda: KNeighborsClassifier(n_neighbors=CFG['KNN_K'], n_jobs=-1, metric='cosine')
    svm_fn = lambda: LinearSVC(C=1.0, max_iter=3000, random_state=CFG['RANDOM_STATE'])

    results = []

    for vec in ['Count Vectorizer', 'Hashing Vectorizer', 'Word2Vec']:
        print(f'  ── KNN + {vec} ──')
        results.append(run_classical_cv('KNN', knn_fn, vec, corpus_np, X_w2v, labels, skf))
        print()

    for vec in ['Count Vectorizer', 'Hashing Vectorizer', 'Word2Vec']:
        print(f'  ── SVM + {vec} ──')
        results.append(run_classical_cv('SVM', svm_fn, vec, corpus_np, X_w2v, labels, skf))
        print()

    if CFG['RUN_LSTM']:
        print('  ── LSTM + Trainable Embeddings ──')
        results.append(run_lstm_cv(tokenized, labels, skf, w2v, variant='trainable'))
        print()
        print('  ── LSTM + Word2Vec Embeddings (frozen) ──')
        results.append(run_lstm_cv(tokenized, labels, skf, w2v, variant='w2v'))
        print()
    else:
        print('  [LSTM skipped — set CFG[\"RUN_LSTM\"] = True to enable]')

    return results


# ════════════════════════════════════════════════════════════════════
#  EXCEL OUTPUT
# ════════════════════════════════════════════════════════════════════
def _style_header(ws) -> None:
    hdr_font = Font(bold=True, color='FFFFFF', size=11)
    hdr_fill = PatternFill(fill_type='solid', fgColor='2F5496')
    for cell in ws[1]:
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for col in ws.columns:
        w = max((len(str(c.value or '')) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(w + 4, 42)


def save_excel(results: list) -> Path:
    print('\n[5/6] Saving results to Excel …')
    path = CFG['RESULTS_DIR'] / 'sentiment_analysis_results.xlsx'

    with pd.ExcelWriter(path, engine='openpyxl') as writer:

        # ── Summary sheet ──────────────────────────────────────────
        rows = []
        for r in results:
            rows.append({
                'Model':                 r['Model'],
                'Vectorizer':            r['Vectorizer'],
                'Accuracy':              round(r['Accuracy_mean'],  4),
                'Accuracy Std':          round(r['Accuracy_std'],   4),
                'Precision':             round(r['Precision_mean'], 4),
                'Precision Std':         round(r['Precision_std'],  4),
                'Recall':                round(r['Recall_mean'],    4),
                'Recall Std':            round(r['Recall_std'],     4),
                'Accuracy  (mean±std)':  f"{r['Accuracy_mean']:.4f} ± {r['Accuracy_std']:.4f}",
                'Precision (mean±std)':  f"{r['Precision_mean']:.4f} ± {r['Precision_std']:.4f}",
                'Recall    (mean±std)':  f"{r['Recall_mean']:.4f} ± {r['Recall_std']:.4f}",
            })

        pd.DataFrame(rows).to_excel(writer, sheet_name='Summary', index=False)
        _style_header(writer.sheets['Summary'])

        # ── Per-fold sheets ────────────────────────────────────────
        for r in results:
            sname = f"{r['Model']}_{r['Vectorizer'][:8]}".replace(' ', '_')[:31]
            fold_df = pd.DataFrame({
                'Fold':      list(range(1, CFG['N_FOLDS'] + 1)),
                'Accuracy':  [round(v, 4) for v in r['Accuracy_folds']],
                'Precision': [round(v, 4) for v in r['Precision_folds']],
                'Recall':    [round(v, 4) for v in r['Recall_folds']],
            })
            stats = pd.DataFrame([
                {'Fold': 'Mean',
                 'Accuracy': round(r['Accuracy_mean'], 4),
                 'Precision': round(r['Precision_mean'], 4),
                 'Recall': round(r['Recall_mean'], 4)},
                {'Fold': 'Std',
                 'Accuracy': round(r['Accuracy_std'], 4),
                 'Precision': round(r['Precision_std'], 4),
                 'Recall': round(r['Recall_std'], 4)},
            ])
            pd.concat([fold_df, stats], ignore_index=True).to_excel(
                writer, sheet_name=sname, index=False
            )
            _style_header(writer.sheets[sname])

    print(f'  Saved: {path}')
    return path


# ════════════════════════════════════════════════════════════════════
#  PLOTS
# ════════════════════════════════════════════════════════════════════
_MODEL_COLOR = {'KNN': '#4C8CBF', 'SVM': '#E07B39', 'LSTM': '#5BAD75'}

def _savefig(fig, name: str) -> None:
    p = CFG['RESULTS_DIR'] / name
    fig.savefig(p, bbox_inches='tight', facecolor='white', dpi=150)
    plt.close(fig)
    print(f'  Saved: {p}')


def save_plots(results: list) -> None:
    print('\n[6/6] Generating plots …')

    plt.rcParams.update({
        'font.family':        'DejaVu Sans',
        'font.size':          10,
        'axes.spines.top':    False,
        'axes.spines.right':  False,
        'axes.grid':          True,
        'grid.alpha':         0.3,
        'figure.facecolor':   'white',
    })

    configs   = [f"{r['Model']}\n{r['Vectorizer']}" for r in results]
    accs      = np.array([r['Accuracy_mean']  for r in results])
    precs     = np.array([r['Precision_mean'] for r in results])
    recs      = np.array([r['Recall_mean']    for r in results])
    acc_stds  = np.array([r['Accuracy_std']   for r in results])
    prec_stds = np.array([r['Precision_std']  for r in results])
    rec_stds  = np.array([r['Recall_std']     for r in results])
    bar_clrs  = [_MODEL_COLOR.get(r['Model'], '#888') for r in results]
    n         = len(results)
    x         = np.arange(n)

    # ── Plot 1: Side-by-side metric bars ────────────────────────────
    fig, axes = plt.subplots(1, 3, figsize=(20, 7))
    fig.suptitle(
        'Sentiment140 Sentiment Analysis — Model × Vectorizer Comparison\n'
        '(10-Fold Stratified CV · 100 K tweets · 50 K positive + 50 K negative)',
        fontsize=12, fontweight='bold', y=1.02,
    )
    for ax, vals, stds, title in zip(
        axes,
        [accs, precs, recs],
        [acc_stds, prec_stds, rec_stds],
        ['Accuracy', 'Precision (weighted)', 'Recall (weighted)'],
    ):
        bars = ax.bar(x, vals, yerr=stds, capsize=5,
                      color=bar_clrs, alpha=0.85, edgecolor='white', linewidth=1.2)
        ax.set_title(title, fontweight='bold', fontsize=11)
        ax.set_xticks(x)
        ax.set_xticklabels(configs, rotation=45, ha='right', fontsize=7.5)
        lo = max(0.0, float(vals.min()) - 0.15)
        hi = min(1.0, float(vals.max()) + 0.08)
        ax.set_ylim(lo, hi)
        ax.set_ylabel('Score')
        for b, v in zip(bars, vals):
            ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.003,
                    f'{v:.3f}', ha='center', va='bottom', fontsize=7, fontweight='bold')

    handles = [mpatches.Patch(color=c, label=m) for m, c in _MODEL_COLOR.items()
               if any(r['Model'] == m for r in results)]
    fig.legend(handles=handles, loc='lower center', ncol=len(handles),
               bbox_to_anchor=(0.5, -0.04), fontsize=10, framealpha=0.9)
    plt.tight_layout()
    _savefig(fig, 'plot1_metrics_comparison.png')

    # ── Plot 2: Heatmap ─────────────────────────────────────────────
    df_heat = pd.DataFrame(
        {'Accuracy': accs, 'Precision': precs, 'Recall': recs},
        index=configs,
    )
    fig, ax = plt.subplots(figsize=(9, max(5, n * 0.7)))
    sns.heatmap(
        df_heat, annot=True, fmt='.4f', cmap='RdYlGn',
        vmin=max(0.5, float(accs.min()) - 0.05),
        vmax=min(1.0, float(accs.max()) + 0.02),
        linewidths=0.6, linecolor='white', ax=ax,
        annot_kws={'size': 9, 'weight': 'bold'},
        cbar_kws={'label': 'Score', 'shrink': 0.8},
    )
    ax.set_title('Performance Heatmap\n(Mean Scores over 10 Folds)',
                 fontsize=12, fontweight='bold', pad=15)
    ax.set_xticklabels(ax.get_xticklabels(), rotation=0, fontsize=10)
    ax.set_yticklabels(ax.get_yticklabels(), rotation=0, fontsize=8)
    plt.tight_layout()
    _savefig(fig, 'plot2_performance_heatmap.png')

    # ── Plot 3: Box plots — fold-level score distribution ───────────
    fold_rows = []
    for r in results:
        cfg_label = f"{r['Model']}\n{r['Vectorizer']}"
        for a, p, rc in zip(r['Accuracy_folds'], r['Precision_folds'], r['Recall_folds']):
            fold_rows += [
                {'Config': cfg_label, 'Metric': 'Accuracy',  'Score': a},
                {'Config': cfg_label, 'Metric': 'Precision', 'Score': p},
                {'Config': cfg_label, 'Metric': 'Recall',    'Score': rc},
            ]
    df_folds = pd.DataFrame(fold_rows)
    unique   = df_folds['Config'].unique()

    fig, axes = plt.subplots(1, 3, figsize=(20, 7))
    fig.suptitle('Fold-Level Score Distribution  (10 Folds per Configuration)',
                 fontsize=12, fontweight='bold')
    for ax, metric in zip(axes, ['Accuracy', 'Precision', 'Recall']):
        sub      = df_folds[df_folds['Metric'] == metric]
        box_data = [sub[sub['Config'] == c]['Score'].values for c in unique]
        bp       = ax.boxplot(box_data, patch_artist=True, notch=False,
                              medianprops={'color': 'black', 'linewidth': 2})
        for patch, r in zip(bp['boxes'], results):
            patch.set_facecolor(_MODEL_COLOR.get(r['Model'], '#888'))
            patch.set_alpha(0.75)
        ax.set_title(metric, fontweight='bold', fontsize=11)
        ax.set_xticks(range(1, len(unique) + 1))
        ax.set_xticklabels(unique, rotation=45, ha='right', fontsize=7.5)
        ymin = max(0.3, float(df_folds['Score'].min()) - 0.05)
        ax.set_ylim(ymin, 1.03)
        ax.set_ylabel('Score')
    handles = [mpatches.Patch(color=c, alpha=0.75, label=m)
               for m, c in _MODEL_COLOR.items() if any(r['Model'] == m for r in results)]
    fig.legend(handles=handles, loc='lower center', ncol=len(handles),
               bbox_to_anchor=(0.5, -0.04), fontsize=10)
    plt.tight_layout()
    _savefig(fig, 'plot3_fold_boxplots.png')

    # ── Plot 4: Ranked grouped bar (sorted by accuracy) ─────────────
    order   = np.argsort(accs)[::-1]
    w       = 0.25
    x_sort  = np.arange(n)
    fig, ax = plt.subplots(figsize=(14, 6))
    ax.bar(x_sort - w, accs[order],  w, label='Accuracy',  color='#4C8CBF', alpha=0.9)
    ax.bar(x_sort,     precs[order], w, label='Precision', color='#E07B39', alpha=0.9)
    ax.bar(x_sort + w, recs[order],  w, label='Recall',    color='#5BAD75', alpha=0.9)
    ax.set_xticks(x_sort)
    ax.set_xticklabels([configs[i] for i in order], rotation=45, ha='right', fontsize=8)
    lo = max(0.0, float(accs.min()) - 0.15)
    hi = min(1.05, float(accs.max()) + 0.07)
    ax.set_ylim(lo, hi)
    ax.set_ylabel('Score')
    ax.set_title('Model Rankings — Sorted by Accuracy (High → Low)',
                 fontweight='bold', fontsize=12)
    ax.legend(fontsize=10)
    plt.tight_layout()
    _savefig(fig, 'plot4_model_ranking.png')

    # ── Plot 5: Radar charts per model ──────────────────────────────
    model_groups = {}
    for r in results:
        model_groups.setdefault(r['Model'], []).append(r)

    nm   = len(model_groups)
    cats = ['Accuracy', 'Precision', 'Recall']
    N    = len(cats)
    ang  = [i / N * 2 * np.pi for i in range(N)] + [0]
    radar_colors = ['#4C8CBF', '#E07B39', '#5BAD75', '#9B59B6']

    fig, axes = plt.subplots(1, nm, figsize=(6 * nm, 6), subplot_kw={'polar': True})
    if nm == 1:
        axes = [axes]

    for ax, (mname, mrs) in zip(axes, model_groups.items()):
        ax.set_theta_offset(np.pi / 2)
        ax.set_theta_direction(-1)
        ax.set_xticks(ang[:-1])
        ax.set_xticklabels(cats, size=10)
        all_vals = [v for r in mrs
                    for v in [r['Accuracy_mean'], r['Precision_mean'], r['Recall_mean']]]
        lo_r = max(0.4, min(all_vals) - 0.05)
        ax.set_ylim(lo_r, 1.0)
        ticks = np.round(np.linspace(lo_r, 1.0, 5), 2)
        ax.set_yticks(ticks)
        ax.set_yticklabels([f'{t:.2f}' for t in ticks], size=7)
        ax.set_title(mname, fontweight='bold', size=13, pad=22)

        for i, r in enumerate(mrs):
            vals = [r['Accuracy_mean'], r['Precision_mean'], r['Recall_mean'], r['Accuracy_mean']]
            col  = radar_colors[i % len(radar_colors)]
            ax.plot(ang, vals, 'o-', lw=2, color=col, label=r['Vectorizer'])
            ax.fill(ang, vals, alpha=0.12, color=col)
        ax.legend(loc='upper right', bbox_to_anchor=(1.45, 1.15), fontsize=8)

    fig.suptitle('Radar Charts — Accuracy / Precision / Recall per Model',
                 fontsize=12, fontweight='bold', y=1.04)
    plt.tight_layout()
    _savefig(fig, 'plot5_radar_charts.png')

    print(f'\n  All plots saved to: {CFG["RESULTS_DIR"].resolve()}/')


# ════════════════════════════════════════════════════════════════════
#  MAIN
# ════════════════════════════════════════════════════════════════════
def main() -> None:
    print('╔══════════════════════════════════════════════════════╗')
    print('║   Sentiment Analysis — Sentiment140 (100 K tweets)  ║')
    print('╚══════════════════════════════════════════════════════╝\n')

    t_start = time.time()

    download_dataset()
    texts, labels = load_data()
    tokenized, corpus_str = preprocess_all(texts)
    w2v     = train_word2vec(tokenized)
    results = train_all(tokenized, corpus_str, labels, w2v)
    save_excel(results)
    save_plots(results)

    elapsed = time.time() - t_start
    print(f'\n  Total runtime: {elapsed/60:.1f} min')

    print('\n╔══════════════════════════════════════════════════════════════╗')
    print('║                     FINAL RESULTS                           ║')
    print('╠══════════════════════════════════════════════════════════════╣')
    print(f"  {'Model':<6}  {'Vectorizer':<26}  {'Accuracy':>9}  {'Precision':>10}  {'Recall':>8}")
    print('  ' + '─' * 64)
    for r in sorted(results, key=lambda x: x['Accuracy_mean'], reverse=True):
        print(
            f"  {r['Model']:<6}  {r['Vectorizer']:<26}  "
            f"{r['Accuracy_mean']:>9.4f}  {r['Precision_mean']:>10.4f}  {r['Recall_mean']:>8.4f}"
        )
    print('╚══════════════════════════════════════════════════════════════╝')
    print(f"\n  Output directory: {CFG['RESULTS_DIR'].resolve()}/")


if __name__ == '__main__':
    main()
